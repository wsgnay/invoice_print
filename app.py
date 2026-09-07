"""Paperdesk: local, stateless document utilities."""
import base64
import io
import json
import os
import zipfile
from contextlib import ExitStack
from pathlib import Path

import pymupdf as fitz
from PIL import Image, ImageOps, UnidentifiedImageError
from docx import Document
from docx.shared import Inches, Pt
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from fastapi.staticfiles import StaticFiles
from starlette.concurrency import run_in_threadpool
from starlette.middleware.trustedhost import TrustedHostMiddleware
from invoice_amount import amount_from_text, recognize_amount

app = FastAPI(title="纸间 Paperdesk")
allowed_hosts = [host.strip() for host in os.getenv("ALLOWED_HOSTS", "*").split(",") if host.strip()]
app.add_middleware(TrustedHostMiddleware, allowed_hosts=allowed_hosts)
MAX_BYTES = 50 * 1024 * 1024
MAX_PAGES = 200


@app.get("/healthz", include_in_schema=False)
def healthz():
    return {"status": "ok"}


def open_document(data: bytes):
    if data.startswith(b"%PDF"):
        doc = fitz.open(stream=data, filetype="pdf")
        if doc.needs_pass:
            doc.close()
            raise ValueError("PDF 已加密，请先解除密码后再导入。")
    else:
        with Image.open(io.BytesIO(data)) as image:
            if image.width * image.height > 40_000_000:
                raise ValueError("图片超过 4000 万像素，请缩小后重试。")
            image = ImageOps.exif_transpose(image)
            if image.mode == "RGBA":
                white = Image.new("RGB", image.size, "white")
                white.paste(image, mask=image.getchannel("A"))
                image = white
            else:
                image = image.convert("RGB")
            out = io.BytesIO()
            image.save(out, format="PNG")
        doc = fitz.open()
        page = doc.new_page(width=image.width, height=image.height)
        page.insert_image(page.rect, stream=out.getvalue())
    if not 0 < len(doc) <= MAX_PAGES:
        doc.close()
        raise ValueError("每个文件需包含 1–200 页。")
    return doc


async def read_upload(file):
    data = await file.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise ValueError("单个文件不能超过 50 MB。")
    if not data:
        raise ValueError("文件为空。")
    return data


@app.post("/api/inspect")
async def inspect(file: UploadFile = File(...), invoice: bool = Form(True)):
    try:
        data = await read_upload(file)
        return await run_in_threadpool(inspect_document, data, invoice)
    except (ValueError, RuntimeError, UnidentifiedImageError, Image.DecompressionBombError) as exc:
        raise HTTPException(400, str(exc)) from exc
    except Exception as exc:
        raise HTTPException(400, "无法读取文件，请确认它是有效的 PDF、JPG、PNG 或 WebP。") from exc


def inspect_document(data, invoice):
    with open_document(data) as doc:
        pages = []
        for page in doc:
            zoom = min(440 / max(page.rect.width, page.rect.height), 1)
            pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
            amount = recognize_amount(page) if invoice else {"amount": None}
            pages.append({"page": page.number, "width": page.rect.width, "height": page.rect.height,
                          "thumbnail": "data:image/png;base64," + base64.b64encode(pix.tobytes("png")).decode(),
                          **amount, "has_text": bool(page.get_text().strip())})
        return {"pages": pages, "is_pdf": data.startswith(b"%PDF")}


def pdf_bytes(doc):
    return doc.tobytes(garbage=4, deflate=True)


@app.post("/api/export")
async def export(files: list[UploadFile] = File(...), spec: str = Form(...)):
    try:
        config = json.loads(spec)
        mode = config.get("mode")
        if mode not in ("invoice", "pdf", "image", "word"):
            raise ValueError("未知操作。")
        refs = config.get("pages", [])
        if not refs or len(refs) > MAX_PAGES:
            raise ValueError("请选择 1–200 页进行导出。")
        with ExitStack() as stack:
            docs = []
            total = 0
            for file in files:
                data = await read_upload(file)
                total += len(data)
                if total > 150 * 1024 * 1024:
                    raise ValueError("本次导出文件合计不能超过 150 MB。")
                docs.append(stack.enter_context(open_document(data)))
            for ref in refs:
                f, p = ref["file"], ref["page"]
                if not isinstance(f, int) or not isinstance(p, int) or not 0 <= f < len(docs) or not 0 <= p < len(docs[f]):
                    raise ValueError("页码无效，请重新导入。")
            if mode == "word":
                word = Document()
                word.styles["Normal"].font.name = "Microsoft YaHei"
                word.styles["Normal"].font.size = Pt(10.5)
                for i, ref in enumerate(refs):
                    if i:
                        word.add_page_break()
                    page = docs[ref["file"]][ref["page"]]
                    text = page.get_text(sort=True).strip()
                    if text:
                        for line in text.splitlines():
                            word.add_paragraph(line)
                    else:
                        pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
                        scale = min(6 / page.rect.width, 8.5 / page.rect.height)
                        word.add_picture(io.BytesIO(pix.tobytes("png")), width=Inches(page.rect.width * scale))
                output = io.BytesIO()
                word.save(output)
                return download(output.getvalue(), "converted.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            result = stack.enter_context(fitz.open())
            if mode in ("invoice", "image"):
                per_sheet = int(config.get("layout", 2 if mode == "invoice" else 1))
                if per_sheet not in (1, 2, 4):
                    raise ValueError("不支持此排版。")
                margin = float(config.get("margin", 12))
                if margin not in (12, 24, 36):
                    raise ValueError("边距无效。")
                landscape = config.get("landscape", False)
                width, height = (842, 595) if landscape else (595, 842)
                cols, rows = (2, 2) if per_sheet == 4 else (1, per_sheet)
                for index, ref in enumerate(refs):
                    if index % per_sheet == 0:
                        sheet = result.new_page(width=width, height=height)
                    slot = index % per_sheet
                    cell_w, cell_h = (width - margin * 2) / cols, (height - margin * 2) / rows
                    x, y = margin + slot % cols * cell_w, margin + slot // cols * cell_h
                    rect = fitz.Rect(x + 6, y + 6, x + cell_w - 6, y + cell_h - 6)
                    source = docs[ref["file"]]
                    if source[ref["page"]].get_contents():
                        sheet.show_pdf_page(rect, source, ref["page"], keep_proportion=True)
            else:
                for ref in refs:
                    result.insert_pdf(docs[ref["file"]], from_page=ref["page"], to_page=ref["page"])
            if mode == "pdf" and config.get("split"):
                buffer = io.BytesIO()
                with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
                    for i in range(len(result)):
                        with fitz.open() as single:
                            single.insert_pdf(result, from_page=i, to_page=i)
                            archive.writestr(f"page-{i+1:03}.pdf", pdf_bytes(single))
                return download(buffer.getvalue(), "split-pages.zip", "application/zip")
            return download(pdf_bytes(result), "paperdesk.pdf", "application/pdf")
    except (ValueError, KeyError, IndexError, TypeError, RuntimeError) as exc:
        raise HTTPException(400, str(exc)) from exc
    except Exception as exc:
        raise HTTPException(400, "转换失败，请检查文件格式，或减少本次处理的页数。") from exc


def download(data, name, media):
    return Response(data, media_type=media, headers={"Content-Disposition": f'attachment; filename="{name}"'})


app.mount("/", StaticFiles(directory=Path(__file__).parent / "static", html=True), name="static")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8765)
