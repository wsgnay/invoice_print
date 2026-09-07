import io
import json
import zipfile

import pymupdf as fitz
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image

from app import app
from invoice_amount import TextBox, amount_from_text, ocr_candidates, positioned_candidates, recognize_amount, upper_rmb

client = TestClient(app)


def test_healthcheck():
    response = client.get('/healthz')
    assert response.status_code == 200
    assert response.json() == {'status': 'ok'}


def sample_pdf():
    with fitz.open() as doc:
        for label in ('FIRST invoice', 'SECOND invoice', 'THIRD invoice'):
            page = doc.new_page(width=500, height=250)
            page.insert_text((30, 60), label)
        return doc.tobytes()


def request_export(mode, pages, **options):
    return client.post('/api/export', files=[('files', ('sample.pdf', sample_pdf(), 'application/pdf'))],
                       data={'spec': json.dumps({'mode': mode, 'pages': [{'file': 0, 'page': p} for p in pages], **options})})


def test_amount_requires_total_label():
    assert amount_from_text('金额 123.00 税额 7.38') is None
    assert amount_from_text('价税合计（小写）￥1,230.45') == '1230.45'
    assert amount_from_text('(小写) -123.45') == '-123.45'
    assert amount_from_text('(小写) 123.45\n(小写) 567.89') is None
    assert amount_from_text('价税合计（大写）壹仟贰佰叁拾元肆角伍分') == '1230.45'
    assert amount_from_text('价税合计 ( 小 写 ) ￥ 9,876.50') == '9876.50'
    assert amount_from_text('电子发票（铁路电子客票）\n票价：￥357.00\n电子客票号：123') == '357.00'
    assert amount_from_text('普通发票\n票价：￥357.00') is None


def test_positioned_pdf_fragments_and_conflicts():
    boxes = [TextBox(10, 10, 60, 20, '价税合计'), TextBox(70, 10, 90, 20, '（小写）'),
             TextBox(100, 10, 150, 20, '￥1,230.45')]
    assert positioned_candidates(boxes) == {'1230.45'}
    assert upper_rmb('壹佰元整') == '100.00'
    assert amount_from_text('价税合计（小写）￥100.00 价税合计（大写）壹佰零壹元整') is None
    assert ocr_candidates('(小写) ￥1.230.45') == {'1230.45'}
    ocr_boxes = [TextBox(10, 10, 60, 20, '(小写)', .74), TextBox(70, 10, 140, 20, '￥1.230.45', .82)]
    assert positioned_candidates(ocr_boxes, confidence=.70, ocr=True) == {'1230.45'}

    railway_boxes = [TextBox(10, 1, 100, 11, '（铁路电子客票）', .73),
                     TextBox(10, 30, 120, 42, '票价：￥357.00', .77)]
    assert positioned_candidates(railway_boxes, confidence=.70, ocr=True) == {'357.00'}


def test_real_pdf_text_objects_are_reassembled():
    with fitz.open() as doc:
        page = doc.new_page()
        page.insert_font(fontname='invoice_cn', fontfile='C:/Windows/Fonts/msyh.ttc')
        page.insert_text((40, 500), '价税合计', fontname='invoice_cn', fontsize=14)
        page.insert_text((130, 500), '（小写）', fontname='invoice_cn', fontsize=14)
        page.insert_text((220, 500), '￥2,345.67', fontname='invoice_cn', fontsize=14)
        recognized = recognize_amount(page, use_ocr=False)
    assert recognized['amount'] == '2345.67'
    assert recognized['amount_source'] == 'pdf_text'


def test_inspect_and_bad_file():
    response = client.post('/api/inspect', files={'file': ('sample.pdf', sample_pdf())})
    assert response.status_code == 200
    assert len(response.json()['pages']) == 3
    assert response.json()['pages'][0]['thumbnail'].startswith('data:image/png;base64,')
    assert client.post('/api/inspect', files={'file': ('bad.pdf', b'broken')}).status_code == 400


def test_invoice_two_up_and_odd_count():
    response = request_export('invoice', [0, 1, 2], layout=2)
    assert response.status_code == 200, response.text
    with fitz.open(stream=response.content, filetype='pdf') as doc:
        assert len(doc) == 2
        assert doc[0].rect == fitz.Rect(0, 0, 595, 842)
        assert doc[0].search_for('FIRST')[0].y1 < 421
        assert doc[0].search_for('SECOND')[0].y0 > 421
        assert 'THIRD' in doc[1].get_text()


def test_reorder_merge_and_split():
    response = request_export('pdf', [2, 0])
    with fitz.open(stream=response.content, filetype='pdf') as doc:
        assert len(doc) == 2
        assert 'THIRD' in doc[0].get_text()
        assert 'FIRST' in doc[1].get_text()
    response = request_export('pdf', [2, 0], split=True)
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert len(archive.namelist()) == 2
        with fitz.open(stream=archive.read('page-001.pdf'), filetype='pdf') as doc:
            assert 'THIRD' in doc[0].get_text()
    response = client.post('/api/export', files=[('files', ('a.pdf', sample_pdf())), ('files', ('b.pdf', sample_pdf()))], data={'spec': json.dumps({'mode':'pdf','pages':[{'file':0,'page':0},{'file':1,'page':2}]})})
    with fitz.open(stream=response.content, filetype='pdf') as doc:
        assert len(doc) == 2 and 'THIRD' in doc[1].get_text()


def test_image_and_word():
    image = Image.new('RGBA', (600, 300), (255, 0, 0, 100))
    data = io.BytesIO()
    image.save(data, format='PNG')
    response = client.post('/api/export', files=[('files', ('image.png', data.getvalue()))], data={'spec':json.dumps({'mode':'image','pages':[{'file':0,'page':0}], 'layout':1})})
    assert response.status_code == 200
    with fitz.open(stream=response.content, filetype='pdf') as doc:
        assert len(doc) == 1
        assert doc[0].get_images()
    response = request_export('word', [1, 0])
    assert response.status_code == 200
    word = Document(io.BytesIO(response.content))
    text = '\n'.join(p.text for p in word.paragraphs)
    assert text.index('SECOND') < text.index('FIRST')
    response = client.post('/api/export', files=[('files', ('image.png', data.getvalue()))], data={'spec':json.dumps({'mode':'word','pages':[{'file':0,'page':0}]})})
    assert len(Document(io.BytesIO(response.content)).inline_shapes) == 1


def test_invalid_page_and_layout():
    assert request_export('pdf', [-1]).status_code == 400
    assert request_export('pdf', [999]).status_code == 400
    assert request_export('invoice', [0], layout=3).status_code == 400
    assert request_export('pdf', []).status_code == 400


def test_encrypted_pdf():
    with fitz.open(stream=sample_pdf(), filetype='pdf') as doc:
        data = doc.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw='owner', user_pw='secret')
    response = client.post('/api/inspect', files={'file': ('locked.pdf', data)})
    assert response.status_code == 400
    assert '加密' in response.json()['detail']
